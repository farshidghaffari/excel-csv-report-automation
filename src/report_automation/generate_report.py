from __future__ import annotations

from pathlib import Path
from typing import Literal

import pandas as pd
from openpyxl.styles import Font, PatternFill


REQUIRED_COLUMNS = {"date", "product", "category", "quantity", "unit_price"}


def _read_input_file(input_path: Path) -> pd.DataFrame:
    """Read a supported source file into a DataFrame."""
    suffix = input_path.suffix.lower()

    if suffix == ".csv":
        return pd.read_csv(input_path)

    if suffix == ".xlsx":
        return pd.read_excel(input_path, engine="openpyxl")

    if suffix == ".xls":
        return pd.read_excel(input_path, engine="xlrd")

    raise ValueError("Unsupported file type. Please provide a .csv, .xlsx, or .xls file.")


def _normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Normalize column names to lowercase snake_case style."""
    df = df.copy()
    df.columns = (
        df.columns.astype(str)
        .str.strip()
        .str.lower()
        .str.replace(" ", "_", regex=False)
        .str.replace("-", "_", regex=False)
    )
    return df


def _validate_columns(df: pd.DataFrame) -> None:
    """Validate that all required columns exist."""
    duplicate_columns = sorted(set(df.columns[df.columns.duplicated()].tolist()))
    if duplicate_columns:
        duplicates = ", ".join(duplicate_columns)
        raise ValueError(f"Duplicate columns after normalization: {duplicates}")

    missing_columns = REQUIRED_COLUMNS - set(df.columns)
    if missing_columns:
        missing = ", ".join(sorted(missing_columns))
        raise ValueError(f"Missing required columns: {missing}")


def _append_rejection_reason(
    reasons: pd.Series,
    mask: pd.Series,
    reason: str,
) -> None:
    """Append a readable reason to each row selected by mask."""
    selected = reasons.loc[mask]
    reasons.loc[mask] = selected.where(selected.eq(""), selected + "; ") + reason


def _clean_and_split_data(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return accepted rows and a traceable rejected-row audit."""
    original_df = df.copy()
    working_df = df.copy()
    source_rows = pd.Series(range(2, len(df) + 2), index=df.index, dtype="int64")

    parsed_dates = pd.to_datetime(working_df["date"], errors="coerce", format="mixed")
    parsed_quantities = pd.to_numeric(working_df["quantity"], errors="coerce")
    parsed_prices = pd.to_numeric(working_df["unit_price"], errors="coerce")
    normalized_products = (
        working_df["product"].astype("string").str.strip().replace("", pd.NA)
    )
    normalized_categories = (
        working_df["category"].astype("string").str.strip().replace("", pd.NA)
    )

    reasons = pd.Series("", index=working_df.index, dtype="string")

    _append_rejection_reason(reasons, parsed_dates.isna(), "invalid or missing date")
    _append_rejection_reason(reasons, normalized_products.isna(), "missing product")
    _append_rejection_reason(reasons, normalized_categories.isna(), "missing category")

    _append_rejection_reason(
        reasons,
        parsed_quantities.isna(),
        "invalid or missing quantity",
    )
    _append_rejection_reason(
        reasons,
        parsed_quantities.notna() & parsed_quantities.le(0),
        "quantity must be greater than zero",
    )
    _append_rejection_reason(
        reasons,
        parsed_quantities.notna() & parsed_quantities.mod(1).ne(0),
        "quantity must be a whole number",
    )

    _append_rejection_reason(
        reasons,
        parsed_prices.isna(),
        "invalid or missing unit price",
    )
    _append_rejection_reason(
        reasons,
        parsed_prices.notna() & parsed_prices.lt(0),
        "unit price cannot be negative",
    )

    accepted_mask = reasons.eq("")

    working_df["date"] = parsed_dates
    working_df["product"] = normalized_products
    working_df["category"] = normalized_categories
    working_df["quantity"] = parsed_quantities
    working_df["unit_price"] = parsed_prices

    clean_df = working_df.loc[accepted_mask].copy()
    clean_df.insert(0, "source_row", source_rows.loc[accepted_mask])
    clean_df["quantity"] = clean_df["quantity"].astype("int64")
    clean_df["total_revenue"] = clean_df["quantity"] * clean_df["unit_price"]
    clean_df["month"] = clean_df["date"].dt.to_period("M").astype(str)
    clean_df = clean_df.sort_values(["date", "source_row"]).reset_index(drop=True)

    rejected_mask = ~accepted_mask
    rejected_df = original_df.loc[rejected_mask].copy()
    rejected_df.insert(0, "source_row", source_rows.loc[rejected_mask])
    rejected_df.insert(1, "rejection_reason", reasons.loc[rejected_mask])
    rejected_df = rejected_df.reset_index(drop=True)

    return clean_df, rejected_df


def _build_summary(
    df: pd.DataFrame,
    group_by: Literal["month", "category", "product"],
) -> pd.DataFrame:
    """Build a grouped revenue summary."""
    summary = df.groupby(group_by, as_index=False, dropna=False).agg(
        total_quantity=("quantity", "sum"),
        total_revenue=("total_revenue", "sum"),
        average_unit_price=("unit_price", "mean"),
        order_count=("product", "count"),
    )

    summary["total_revenue"] = summary["total_revenue"].round(2)
    summary["average_unit_price"] = summary["average_unit_price"].round(2)

    if group_by == "month":
        return summary.sort_values("month").reset_index(drop=True)

    return summary.sort_values("total_revenue", ascending=False).reset_index(drop=True)


def _format_workbook(writer: pd.ExcelWriter) -> None:
    """Apply lightweight, business-readable formatting to every worksheet."""
    header_fill = PatternFill(fill_type="solid", fgColor="12372A")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font

        if worksheet.max_row > 1:
            worksheet.auto_filter.ref = worksheet.dimensions

        for column_cells in worksheet.columns:
            values = ["" if cell.value is None else str(cell.value) for cell in column_cells]
            width = min(max(max(map(len, values), default=0) + 2, 12), 48)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width

    cleaned_sheet = writer.book["Cleaned Data"]
    cleaned_headers = {cell.value: cell.column for cell in cleaned_sheet[1]}
    for header in ("unit_price", "total_revenue"):
        column = cleaned_headers.get(header)
        if column:
            for row in range(2, cleaned_sheet.max_row + 1):
                cleaned_sheet.cell(row=row, column=column).number_format = "#,##0.00"

    for sheet_name in ("Monthly Summary", "Category Summary", "Product Summary"):
        worksheet = writer.book[sheet_name]
        headers = {cell.value: cell.column for cell in worksheet[1]}
        for header in ("total_revenue", "average_unit_price"):
            column = headers.get(header)
            if column:
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row, column=column).number_format = "#,##0.00"


def generate_report(input_file: str | Path, output_file: str | Path) -> Path:
    """Generate a traceable Excel report from a supported source file."""
    input_path = Path(input_file)
    output_path = Path(output_file)

    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("Output file must use the .xlsx extension.")

    output_path.parent.mkdir(parents=True, exist_ok=True)

    raw_df = _read_input_file(input_path)
    normalized_df = _normalize_columns(raw_df)
    _validate_columns(normalized_df)

    clean_df, rejected_df = _clean_and_split_data(normalized_df)

    monthly_summary = _build_summary(clean_df, "month")
    category_summary = _build_summary(clean_df, "category")
    product_summary = _build_summary(clean_df, "product")

    report_info = pd.DataFrame(
        {
            "metric": [
                "source_file",
                "input_rows",
                "accepted_rows",
                "rejected_rows",
                "total_revenue",
                "date_from",
                "date_to",
            ],
            "value": [
                str(input_path),
                len(normalized_df),
                len(clean_df),
                len(rejected_df),
                round(clean_df["total_revenue"].sum(), 2),
                clean_df["date"].min().strftime("%Y-%m-%d") if not clean_df.empty else "",
                clean_df["date"].max().strftime("%Y-%m-%d") if not clean_df.empty else "",
            ],
        }
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        clean_df.to_excel(writer, sheet_name="Cleaned Data", index=False)
        rejected_df.to_excel(writer, sheet_name="Rejected Rows", index=False)
        monthly_summary.to_excel(writer, sheet_name="Monthly Summary", index=False)
        category_summary.to_excel(writer, sheet_name="Category Summary", index=False)
        product_summary.to_excel(writer, sheet_name="Product Summary", index=False)
        report_info.to_excel(writer, sheet_name="Report Info", index=False)
        _format_workbook(writer)

    return output_path
