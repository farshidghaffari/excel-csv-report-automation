from pathlib import Path
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from report_automation.generate_report import generate_report


VALID_HEADER = "date,product,category,quantity,unit_price\n"


def test_generate_report_creates_structured_workbook(tmp_path: Path) -> None:
    input_file = tmp_path / "input.csv"
    output_file = tmp_path / "report.xlsx"

    input_file.write_text(
        VALID_HEADER
        + "2026-01-01,USB Cable,Accessories,2,8.5\n"
        + "2026-02-02,Keyboard,Hardware,1,45\n",
        encoding="utf-8",
    )

    result = generate_report(input_file=input_file, output_file=output_file)

    assert result == output_file
    assert result.exists()

    workbook = load_workbook(result)
    assert workbook.sheetnames == [
        "Cleaned Data",
        "Rejected Rows",
        "Monthly Summary",
        "Category Summary",
        "Product Summary",
        "Report Info",
    ]
    assert all(sheet.freeze_panes == "A2" for sheet in workbook.worksheets)


def test_generate_report_calculates_revenue_and_report_metrics(tmp_path: Path) -> None:
    input_file = tmp_path / "input.csv"
    output_file = tmp_path / "report.xlsx"

    input_file.write_text(
        VALID_HEADER
        + "2026-01-01,USB Cable,Accessories,2,8.5\n"
        + "2026-01-02,Keyboard,Hardware,1,45\n",
        encoding="utf-8",
    )

    generate_report(input_file, output_file)

    cleaned = pd.read_excel(output_file, sheet_name="Cleaned Data")
    report_info = pd.read_excel(output_file, sheet_name="Report Info")
    metrics = dict(zip(report_info["metric"], report_info["value"]))

    assert cleaned["total_revenue"].tolist() == [17.0, 45.0]
    assert metrics["input_rows"] == 2
    assert metrics["accepted_rows"] == 2
    assert metrics["rejected_rows"] == 0
    assert metrics["total_revenue"] == 62.0


def test_invalid_rows_are_audited_instead_of_silently_dropped(tmp_path: Path) -> None:
    input_file = tmp_path / "mixed_input.csv"
    output_file = tmp_path / "report.xlsx"

    input_file.write_text(
        VALID_HEADER
        + "2026-01-01,USB Cable,Accessories,2,8.5\n"
        + "not-a-date,Mouse,Accessories,1,20\n"
        + "2026-01-03,Keyboard,Hardware,not-a-number,45\n"
        + "2026-01-04,Monitor,Hardware,1,-10\n"
        + "2026-01-05,,Hardware,1,25\n",
        encoding="utf-8",
    )

    generate_report(input_file, output_file)

    cleaned = pd.read_excel(output_file, sheet_name="Cleaned Data")
    rejected = pd.read_excel(output_file, sheet_name="Rejected Rows")
    report_info = pd.read_excel(output_file, sheet_name="Report Info")
    metrics = dict(zip(report_info["metric"], report_info["value"]))

    assert len(cleaned) == 1
    assert len(rejected) == 4
    assert rejected["source_row"].tolist() == [3, 4, 5, 6]
    assert "invalid or missing date" in rejected.loc[0, "rejection_reason"]
    assert "invalid or missing quantity" in rejected.loc[1, "rejection_reason"]
    assert "unit price cannot be negative" in rejected.loc[2, "rejection_reason"]
    assert "missing product" in rejected.loc[3, "rejection_reason"]
    assert metrics["accepted_rows"] == 1
    assert metrics["rejected_rows"] == 4


def test_generate_report_accepts_xlsx_input(tmp_path: Path) -> None:
    input_file = tmp_path / "input.xlsx"
    output_file = tmp_path / "report.xlsx"

    pd.DataFrame(
        [
            {
                "date": "2026-01-01",
                "product": "USB Cable",
                "category": "Accessories",
                "quantity": 2,
                "unit_price": 8.5,
            }
        ]
    ).to_excel(input_file, index=False)

    result = generate_report(input_file, output_file)

    assert result.exists()


def test_generate_report_rejects_missing_columns(tmp_path: Path) -> None:
    input_file = tmp_path / "bad_input.csv"
    output_file = tmp_path / "report.xlsx"

    input_file.write_text(
        "date,product,quantity\n"
        "2026-01-01,USB Cable,2\n",
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="Missing required columns"):
        generate_report(input_file=input_file, output_file=output_file)


def test_generate_report_rejects_columns_that_collide_after_normalization(
    tmp_path: Path,
) -> None:
    input_file = tmp_path / "duplicate_columns.xlsx"
    output_file = tmp_path / "report.xlsx"

    pd.DataFrame(
        [["2026-01-01", "USB Cable", "Accessories", 2, 8.5, 9.5]],
        columns=[
            "date",
            "product",
            "category",
            "quantity",
            "Unit Price",
            "unit-price",
        ],
    ).to_excel(input_file, index=False)

    with pytest.raises(ValueError, match="Duplicate columns after normalization: unit_price"):
        generate_report(input_file, output_file)


def test_generate_report_rejects_unsupported_input_type(tmp_path: Path) -> None:
    input_file = tmp_path / "input.json"
    output_file = tmp_path / "report.xlsx"
    input_file.write_text("{}", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file type"):
        generate_report(input_file, output_file)


def test_generate_report_requires_xlsx_output(tmp_path: Path) -> None:
    input_file = tmp_path / "input.csv"
    output_file = tmp_path / "report.csv"
    input_file.write_text(
        VALID_HEADER + "2026-01-01,USB Cable,Accessories,2,8.5\n",
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="Output file must use the .xlsx extension"):
        generate_report(input_file, output_file)
